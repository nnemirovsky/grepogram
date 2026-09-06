"""The extractor registry, the document dispatcher, and ``FakeClient``'s fetch/download pair.

The committed fixtures (``tests/fixtures/sample.pdf``, ``tests/fixtures/sample.docx``) are the
round-trip cases; the ``_write_*`` helpers build the oversized and malformed files a fixture
should not be. Nothing here downloads anything.
"""

import importlib.util
import logging
import sys
import zipfile
from pathlib import Path

import pytest

from grepogram import extract
from grepogram.extract import EXTRACT_MAX_CHARS, ExtractError
from tests.fakes import FakeClient, make_message

FIXTURES = Path(__file__).resolve().parent / "fixtures"
SAMPLE_PDF = FIXTURES / "sample.pdf"
SAMPLE_DOCX = FIXTURES / "sample.docx"


def _needs(library: str) -> None:
    """Skip a test that parses a real file where its optional library is not installed.

    The ``media`` extra is optional and CI installs no extras, so what this module must never do
    is *fail* there: the "no extractor here" path is the product's documented degradation and has
    tests of its own, all of which monkeypatch and run anywhere.
    """
    pytest.importorskip(library, reason=f"{library} is not installed")


def _write_pdf(path: Path, lines: list[str]) -> Path:
    """A one-page PDF drawing ``lines``, written by hand so no generator library is needed.

    Reading one back still needs ``pypdf``, so every caller is a test that parses it."""
    _needs("pypdf")
    drawn = b"".join(b"(" + line.encode("ascii") + b") Tj 0 -20 Td " for line in lines)
    stream = b"BT /F1 14 Tf 72 720 Td " + drawn + b"ET\n"
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        b"<< /Length " + str(len(stream)).encode() + b" >>\nstream\n" + stream + b"endstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n" % (len(objects) + 1)
    out += b"0000000000 65535 f \n"
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\n" % (len(objects) + 1)
    out += b"startxref\n%d\n%%%%EOF\n" % xref
    path.write_bytes(bytes(out))
    return path


def _write_docx(path: Path, paragraphs: list[str], table: list[list[str]] | None = None) -> Path:
    """Build a DOCX with ``python-docx``, skipping the test where the library is not installed.

    Imported here rather than at module scope: the ``media`` extra is optional and CI installs
    no extras, so a top-level import would make the whole module uncollectable — including the
    tests of the degradation path that exists for exactly that installation.
    """
    _needs("docx")
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        built = document.add_table(rows=len(table), cols=len(table[0]))
        for row, cells in zip(built.rows, table, strict=True):
            for cell, text in zip(row.cells, cells, strict=True):
                cell.text = text
    document.save(str(path))
    return path


# --- the committed fixtures ------------------------------------------------------------------


def test_pdf_fixture_round_trips() -> None:
    _needs("pypdf")
    text = extract.extract_document(SAMPLE_PDF)
    assert text.splitlines() == ["Grepogram sample PDF", "Embassy notice 2026"]


def test_docx_fixture_round_trips() -> None:
    _needs("docx")
    text = extract.extract_document(SAMPLE_DOCX)
    assert text.splitlines() == ["Grepogram sample DOCX", "Rental contract clause"]


def test_fixtures_stay_small() -> None:
    """Committed binaries, so they are hand-built rather than downloaded — keep them tiny."""
    assert SAMPLE_PDF.stat().st_size < 4096
    assert SAMPLE_DOCX.stat().st_size < 4096


# --- dispatch --------------------------------------------------------------------------------


def test_docx_extension_over_pdf_bytes_is_rejected(tmp_path: Path) -> None:
    mislabelled = tmp_path / "notice.docx"
    mislabelled.write_bytes(SAMPLE_PDF.read_bytes())
    with pytest.raises(ExtractError, match="does not hold docx data"):
        extract.extract_document(mislabelled)


def test_pdf_extension_over_docx_bytes_is_rejected(tmp_path: Path) -> None:
    mislabelled = tmp_path / "contract.pdf"
    mislabelled.write_bytes(SAMPLE_DOCX.read_bytes())
    with pytest.raises(ExtractError, match="does not hold pdf data"):
        extract.extract_document(mislabelled)


def test_dispatch_ignores_extension_case(tmp_path: Path) -> None:
    _needs("pypdf")
    shouted = tmp_path / "NOTICE.PDF"
    shouted.write_bytes(SAMPLE_PDF.read_bytes())
    assert "Embassy notice 2026" in extract.extract_document(shouted)


def test_unknown_extension_has_no_extractor(tmp_path: Path) -> None:
    other = tmp_path / "prices.xlsx"
    other.write_bytes(b"PK\x03\x04whatever")
    with pytest.raises(ExtractError, match="no document extractor for .xlsx"):
        extract.extract_document(other)


def test_a_name_with_no_extension_has_no_extractor(tmp_path: Path) -> None:
    nameless = tmp_path / "attachment"
    nameless.write_bytes(SAMPLE_PDF.read_bytes())
    with pytest.raises(ExtractError, match="no document extractor for a name with no extension"):
        extract.extract_document(nameless)


def test_a_missing_file_is_an_extract_error(tmp_path: Path) -> None:
    _needs("pypdf")
    with pytest.raises(ExtractError, match="cannot read gone.pdf"):
        extract.extract_document(tmp_path / "gone.pdf")


# --- failures --------------------------------------------------------------------------------


def test_corrupt_pdf_is_an_extract_error(tmp_path: Path) -> None:
    _needs("pypdf")
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"%PDF-1.4\nnot really a pdf at all\n")
    with pytest.raises(ExtractError, match="cannot read broken.pdf as a PDF"):
        extract.extract_document(broken)


def test_corrupt_docx_is_an_extract_error(tmp_path: Path) -> None:
    _needs("docx")
    broken = tmp_path / "broken.docx"
    broken.write_bytes(b"PK\x03\x04not a real zip")
    with pytest.raises(ExtractError, match="cannot read broken.docx as a DOCX"):
        extract.extract_document(broken)


def test_a_zip_that_is_not_a_docx_is_an_extract_error(tmp_path: Path) -> None:
    """Magic bytes only prove it is a zip; ``python-docx`` is what proves it is a document."""
    _needs("docx")
    not_a_document = tmp_path / "photos.docx"
    with zipfile.ZipFile(not_a_document, "w") as archive:
        archive.writestr("hello.txt", "no OOXML parts here")
    with pytest.raises(ExtractError, match="cannot read photos.docx as a DOCX"):
        extract.extract_document(not_a_document)


def test_pdf_without_pypdf_is_an_extract_error(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(sys.modules, "pypdf", None)
    with pytest.raises(ExtractError, match="pypdf is not installed"):
        extract.extract_pdf(SAMPLE_PDF)


def test_docx_without_python_docx_is_an_extract_error(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setitem(sys.modules, "docx", None)
    with pytest.raises(ExtractError, match="python-docx is not installed"):
        extract.extract_docx(SAMPLE_DOCX)


# --- what the extractors read ------------------------------------------------------------------


def test_docx_reads_tables_too(tmp_path: Path) -> None:
    path = _write_docx(
        tmp_path / "prices.docx", ["Price list"], table=[["Visa", "120 EUR"], ["NIE", "12 EUR"]]
    )
    text = extract.extract_document(path)
    assert text.splitlines() == ["Price list", "Visa", "120 EUR", "NIE", "12 EUR"]


def test_blank_lines_are_dropped(tmp_path: Path) -> None:
    path = _write_docx(tmp_path / "sparse.docx", ["First", "", "   ", "Second"])
    assert extract.extract_document(path).splitlines() == ["First", "Second"]


def test_an_empty_document_extracts_to_nothing(tmp_path: Path) -> None:
    path = _write_docx(tmp_path / "empty.docx", [])
    assert extract.extract_document(path) == ""


# --- the length cap ----------------------------------------------------------------------------


def test_a_long_docx_is_capped(tmp_path: Path) -> None:
    path = _write_docx(tmp_path / "long.docx", ["x" * 100 for _ in range(100)])
    text = extract.extract_document(path)
    assert len(text) == EXTRACT_MAX_CHARS


def test_a_long_pdf_is_capped(tmp_path: Path) -> None:
    path = _write_pdf(tmp_path / "long.pdf", ["y" * 80 for _ in range(100)])
    text = extract.extract_document(path)
    assert len(text) == EXTRACT_MAX_CHARS


def test_text_at_the_cap_is_kept_whole(tmp_path: Path) -> None:
    line = "z" * EXTRACT_MAX_CHARS
    path = _write_docx(tmp_path / "exact.docx", [line])
    assert extract.extract_document(path) == line


# --- macOS OCR ---------------------------------------------------------------------------------

# The stand-ins below are shaped after the real framework, which was driven once by hand on
# macOS 26 before they were written: ``alloc().init()``, the two-value ``(value, error)`` returns
# pyobjc makes of an ``NSError **`` out-parameter, ``performRequests:error:`` taking an array, and
# ``topCandidates_(1)`` handing back objects whose ``string()`` is the recognised line. CI has no
# Vision, so replacing ``extract._vision`` with one of these is how the extractor is exercised.


class _FakeText:
    def __init__(self, text: str) -> None:
        self._text = text

    def string(self) -> str:
        return self._text


class _FakeObservation:
    def __init__(self, *candidates: str) -> None:
        self._candidates = [_FakeText(text) for text in candidates]

    def topCandidates_(self, count: int) -> list[_FakeText]:  # a pyobjc selector name
        return self._candidates[:count]


class _FakeRequest:
    def __init__(self, vision: "_FakeVision") -> None:
        self._vision = vision
        self.level: object = None
        self.correction: bool | None = None
        self.languages: list[str] = []

    def setRecognitionLevel_(self, level: object) -> None:  # a pyobjc selector name
        self.level = level

    def setUsesLanguageCorrection_(self, on: bool) -> None:  # a pyobjc selector name
        self.correction = on

    def supportedRecognitionLanguagesAndReturnError_(  # a pyobjc selector name
        self, error: None
    ) -> tuple[list[str], object]:
        return self._vision.supported, self._vision.languages_error

    def setRecognitionLanguages_(self, languages: list[str]) -> None:  # a pyobjc selector name
        self.languages = list(languages)

    def results(self) -> list[_FakeObservation] | None:
        return self._vision.results


class _RequestClass:
    """``Vision.VNRecognizeTextRequest``: ``alloc().init()`` hands back the one request."""

    def __init__(self, vision: "_FakeVision") -> None:
        self._vision = vision

    def alloc(self) -> "_RequestClass":
        return self

    def init(self) -> _FakeRequest:
        self._vision.request = _FakeRequest(self._vision)
        return self._vision.request


class _FakeHandler:
    def __init__(self, vision: "_FakeVision") -> None:
        self._vision = vision

    def performRequests_error_(  # a pyobjc selector name
        self, requests: list[_FakeRequest], error: None
    ) -> tuple[bool, object]:
        self._vision.performed = list(requests)
        return self._vision.done, self._vision.perform_error


class _HandlerClass:
    """``Vision.VNImageRequestHandler``: built from the image bytes, which it records."""

    def __init__(self, vision: "_FakeVision") -> None:
        self._vision = vision

    def alloc(self) -> "_HandlerClass":
        return self

    def initWithData_options_(  # a pyobjc selector name
        self, data: bytes, options: dict[str, object]
    ) -> _FakeHandler:
        self._vision.image = data
        self._vision.options = options
        return _FakeHandler(self._vision)


class _FakeVision:
    """The ``Vision`` module as ``extract._recognise`` uses it, and a log of what it was asked."""

    VNRequestTextRecognitionLevelAccurate = "accurate"

    def __init__(
        self,
        *,
        lines: tuple[str, ...] = (),
        supported: tuple[str, ...] = ("ru-RU", "en-US", "tr-TR"),
        languages_error: object = None,
        done: bool = True,
        perform_error: object = None,
    ) -> None:
        self.supported = list(supported)
        self.languages_error = languages_error
        self.done = done
        self.perform_error = perform_error
        self.results: list[_FakeObservation] | None = [_FakeObservation(line) for line in lines]
        self.request: _FakeRequest | None = None
        self.image: bytes | None = None
        self.options: dict[str, object] | None = None
        self.performed: list[_FakeRequest] = []
        self.VNRecognizeTextRequest = _RequestClass(self)
        self.VNImageRequestHandler = _HandlerClass(self)


def _photo(tmp_path: Path, name: str = "notice.jpg") -> Path:
    path = tmp_path / name
    path.write_bytes(b"\xff\xd8\xff\xe0 jpeg bytes")
    return path


def _fake_vision(monkeypatch: pytest.MonkeyPatch, vision: _FakeVision) -> _FakeVision:
    monkeypatch.setattr(extract, "_vision", lambda: vision)
    return vision


def test_ocr_returns_the_lines_vision_recognised(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    vision = _fake_vision(
        monkeypatch, _FakeVision(lines=("Посольство Испании", "Embassy notice 2026"))
    )
    photo = _photo(tmp_path)
    assert extract.ocr_image(photo) == "Посольство Испании\nEmbassy notice 2026"
    assert vision.image == photo.read_bytes()
    assert vision.performed == [vision.request]


def test_ocr_asks_for_the_accurate_level_and_language_correction(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    vision = _fake_vision(monkeypatch, _FakeVision(lines=("text",)))
    extract.ocr_image(_photo(tmp_path))
    assert vision.request is not None
    assert vision.request.level == _FakeVision.VNRequestTextRecognitionLevelAccurate
    assert vision.request.correction is True


def test_ocr_requests_the_languages_this_build_supports(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    vision = _fake_vision(monkeypatch, _FakeVision(lines=("text",)))
    extract.ocr_image(_photo(tmp_path))
    assert vision.request is not None
    assert vision.request.languages == ["ru-RU", "en-US", "tr-TR"]


def test_ocr_drops_a_language_the_build_does_not_offer(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """macOS 14 has no Russian, and a request asking for it fails outright — so it is not asked."""
    vision = _fake_vision(monkeypatch, _FakeVision(lines=("notice",), supported=("en-US", "fr-FR")))
    assert extract.ocr_image(_photo(tmp_path)) == "notice"
    assert vision.request is not None
    assert vision.request.languages == ["en-US"]


def test_a_photo_with_no_text_extracts_to_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_vision(monkeypatch, _FakeVision())
    assert extract.ocr_image(_photo(tmp_path)) == ""


def test_no_results_at_all_extracts_to_nothing(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Vision answers ``nil`` rather than an empty array when it recognised nothing."""
    vision = _fake_vision(monkeypatch, _FakeVision())
    vision.results = None
    assert extract.ocr_image(_photo(tmp_path)) == ""


def test_an_observation_with_no_candidate_is_skipped(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    vision = _fake_vision(monkeypatch, _FakeVision(lines=("kept",)))
    vision.results = [_FakeObservation(), _FakeObservation("kept")]
    assert extract.ocr_image(_photo(tmp_path)) == "kept"


def test_ocr_text_is_capped(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    _fake_vision(monkeypatch, _FakeVision(lines=tuple("x" * 100 for _ in range(100))))
    assert len(extract.ocr_image(_photo(tmp_path))) == EXTRACT_MAX_CHARS


def test_a_failed_vision_request_is_an_extract_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_vision(monkeypatch, _FakeVision(done=False, perform_error="zero-dimensioned image"))
    with pytest.raises(ExtractError, match="Vision could not read the image"):
        extract.ocr_image(_photo(tmp_path))


def test_a_language_query_error_is_an_extract_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_vision(monkeypatch, _FakeVision(languages_error="Code=1"))
    with pytest.raises(ExtractError, match="cannot report its recognition languages"):
        extract.ocr_image(_photo(tmp_path))


def test_an_unexpected_framework_error_is_an_extract_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_vision(monkeypatch, _FakeVision())

    def boom(vision: object, image: bytes) -> list[str]:
        raise RuntimeError("objc: unrecognised selector")

    monkeypatch.setattr(extract, "_recognise", boom)
    with pytest.raises(ExtractError, match="cannot read text from notice.jpg"):
        extract.ocr_image(_photo(tmp_path))


def test_ocr_of_a_missing_file_is_an_extract_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    _fake_vision(monkeypatch, _FakeVision())
    with pytest.raises(ExtractError, match="cannot read text from gone.jpg"):
        extract.ocr_image(tmp_path / "gone.jpg")


def test_ocr_without_pyobjc_is_an_extract_error(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    def missing() -> object:
        raise ImportError("No module named 'Vision'")

    monkeypatch.setattr(extract, "_vision", missing)
    with pytest.raises(ExtractError, match="pyobjc-framework-Vision is not installed"):
        extract.ocr_image(_photo(tmp_path))


def test_the_vision_seam_hands_back_the_framework(monkeypatch: pytest.MonkeyPatch) -> None:
    """The one call CI cannot make; ``sys.modules`` is what stands in for the framework here."""
    stand_in = _FakeVision()
    monkeypatch.setitem(sys.modules, "Vision", stand_in)
    assert extract._vision() is stand_in


def test_the_vision_seam_raises_import_error_without_the_framework(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setitem(sys.modules, "Vision", None)
    with pytest.raises(ImportError):
        extract._vision()


# --- the recognition languages -----------------------------------------------------------------


def test_supported_languages_are_asked_for_best_first() -> None:
    assert extract._requested_languages(["en-US", "fr-FR", "ru-RU"]) == ["ru-RU", "en-US"]


def test_an_unsupported_language_is_dropped_rather_than_requested() -> None:
    assert extract._requested_languages(["en-US", "fr-FR"]) == ["en-US"]


def test_a_regional_variant_stands_in_for_the_wanted_tag() -> None:
    assert extract._requested_languages(["en-GB", "ru-RU"]) == ["ru-RU", "en-GB"]


def test_the_whole_tag_wins_over_a_shared_primary_subtag() -> None:
    assert extract._requested_languages(["en-GB", "en-US"]) == ["en-US"]


def test_a_build_offering_neither_language_requests_nothing() -> None:
    """An empty list is Vision's "use your default" — better than a request that fails."""
    assert extract._requested_languages(["fr-FR", "zh-Hans"]) == []


def test_one_offer_is_never_requested_twice(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "OCR_LANGUAGES", ("en-US", "en-GB"))
    assert extract._requested_languages(["en-US"]) == ["en-US"]


# --- whether OCR can run at all ------------------------------------------------------------------


def test_ocr_is_unavailable_off_darwin(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "platform", "linux")
    assert extract._ocr_unavailable() == "macOS Vision needs darwin, not linux"


def test_ocr_is_unavailable_without_the_extra(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "platform", "darwin")
    monkeypatch.setattr(importlib.util, "find_spec", lambda name: None)
    assert extract._ocr_unavailable() == "the 'media' extra brings pyobjc-framework-Vision"


def test_ocr_is_available_on_a_mac_carrying_the_extra(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(sys, "platform", "darwin")
    monkeypatch.setattr(importlib.util, "find_spec", lambda name: object())
    assert extract._ocr_unavailable() is None


class _StubPage:
    """One PDF page that records the moment it was asked for its text."""

    def __init__(self, number: int, read: list[int]) -> None:
        self.number = number
        self._read = read

    def extract_text(self) -> str:
        self._read.append(self.number)
        return "z" * 3000


def test_a_long_pdf_stops_being_parsed_past_the_cap(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """Extraction of a 400-page contract must cost the pages it needs, not the whole document.

    Driven through a stub ``pypdf`` rather than a fixture: the point is *which pages were read*,
    which no committed file can show, and the stub also lets this run where pypdf is absent.
    """
    read: list[int] = []

    class _StubReader:
        def __init__(self, path: str) -> None:
            self.pages = [_StubPage(number, read) for number in range(20)]

    class _StubPypdf:
        PdfReader = _StubReader

    monkeypatch.setitem(sys.modules, "pypdf", _StubPypdf)
    path = tmp_path / "contract.pdf"
    path.write_bytes(b"%PDF-1.4\n")
    assert len(extract.extract_document(path)) == EXTRACT_MAX_CHARS
    assert read == [0, 1], "the cap was passed on the second page and the rest was never parsed"


def test_the_cap_does_not_leave_a_dangling_space(tmp_path: Path) -> None:
    """The cut lands wherever the character count runs out, and a truncation ending in a space
    (or in a newline, half a line into the next one) is text with a ragged edge — one that shows
    up in a unit's rendered line and in a search snippet."""
    assert extract._capped("x " * EXTRACT_MAX_CHARS) == "x " * (EXTRACT_MAX_CHARS // 2 - 1) + "x"
    assert extract._capped("a" * (EXTRACT_MAX_CHARS - 1) + " b").endswith("a")


def test_the_cap_leaves_a_short_text_exactly_as_it_is() -> None:
    assert extract._capped("  first \n\n  second  ") == "first\nsecond"


# --- the registry ------------------------------------------------------------------------------


def test_registry_maps_document(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "_documents_available", lambda: True)
    assert extract.registry()["document"] is extract.extract_document


def test_registry_leaves_the_kinds_with_no_extractor_out() -> None:
    """An absent kind is marked unsupported by the pass; voice and video notes wait for v0.3.0."""
    built = extract.registry()
    assert "voice" not in built
    assert "video_note" not in built
    assert "video" not in built


def test_registry_maps_photo_where_vision_can_run(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(extract, "_ocr_unavailable", lambda: None)
    assert extract.registry()["photo"] is extract.ocr_image


def test_registry_drops_photo_where_vision_cannot(
    monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture
) -> None:
    """Unmapped, so the pass parks photos as unsupported — and one debug line says why."""
    monkeypatch.setattr(extract, "_ocr_unavailable", lambda: "macOS Vision needs darwin, not linux")
    with caplog.at_level(logging.DEBUG, logger="grepogram.extract"):
        built = extract.registry()
    assert "photo" not in built
    assert "no OCR extractor: macOS Vision needs darwin, not linux" in caplog.text


def test_registry_drops_documents_when_the_libraries_are_absent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    monkeypatch.setattr(extract, "_documents_available", lambda: False)
    monkeypatch.setattr(extract, "_ocr_unavailable", lambda: "no Vision here")
    assert extract.registry() == {}


def test_registry_is_derived_on_every_call(monkeypatch: pytest.MonkeyPatch) -> None:
    """An import-time constant could not answer differently after the environment changes."""
    monkeypatch.setattr(extract, "_documents_available", lambda: True)
    assert "document" in extract.registry()
    monkeypatch.setattr(extract, "_documents_available", lambda: False)
    assert "document" not in extract.registry()


def test_documents_are_available_when_either_library_imports(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    """Read off ``find_spec``, never off this machine: the answer has to be asserted for an
    installation without the ``media`` extra too, which is the one that degrades."""
    monkeypatch.setattr(importlib.util, "find_spec", lambda name: object())
    assert extract._documents_available() is True
    monkeypatch.setattr(
        importlib.util, "find_spec", lambda name: None if name == "pypdf" else object()
    )
    assert extract._documents_available() is True
    monkeypatch.setattr(
        importlib.util, "find_spec", lambda name: None if name == "docx" else object()
    )
    assert extract._documents_available() is True
    monkeypatch.setattr(importlib.util, "find_spec", lambda name: None)
    assert extract._documents_available() is False


# --- FakeClient's fetch and download -----------------------------------------------------------

CHAT_ID = -1001234567890

# The three ``get_messages`` tests below are about the *shape* of an answer, not about resolving
# the peer that was asked for, so they switch the entity cache off: a real client would have
# listed its dialogs first (``sync.warm_peer_cache``), and saying so here would only add a dialog
# fixture to assertions that never look at one. The resolution rule has its own tests in
# ``tests/test_tg.py``.


async def test_fake_get_messages_answers_none_for_a_missing_id() -> None:
    client = FakeClient(
        messages={CHAT_ID: [make_message(CHAT_ID, 1, "one")]}, strict_entities=False
    )
    got = await client.get_messages(CHAT_ID, ids=[1, 2])
    assert [None if m is None else m.id for m in got] == [1, None]


async def test_fake_get_messages_answers_one_message_for_an_int_id() -> None:
    client = FakeClient(
        messages={CHAT_ID: [make_message(CHAT_ID, 7, "seven")]}, strict_entities=False
    )
    one = await client.get_messages(CHAT_ID, ids=7)
    assert one.id == 7
    assert await client.get_messages(CHAT_ID, ids=8) is None


async def test_fake_get_messages_defaults_to_one_message() -> None:
    client = FakeClient(
        messages={CHAT_ID: [make_message(CHAT_ID, 1, "one"), make_message(CHAT_ID, 2, "two")]},
        strict_entities=False,
    )
    got = await client.get_messages(CHAT_ID)
    assert [m.id for m in got] == [2]
    assert ("get_messages", {"chat_id": CHAT_ID, "limit": None, "ids": None}) in client.calls


async def test_fake_download_media_writes_the_registered_bytes(tmp_path: Path) -> None:
    message = make_message(CHAT_ID, 3, "photo")
    client = FakeClient(downloads={(CHAT_ID, 3): b"%PDF-bytes"})
    target = tmp_path / "3.pdf"
    assert await client.download_media(message, file=target) == str(target)
    assert target.read_bytes() == b"%PDF-bytes"
    assert ("download_media", {"chat_id": CHAT_ID, "msg_id": 3, "file": target}) in client.calls


async def test_fake_download_media_is_none_when_nothing_is_registered(tmp_path: Path) -> None:
    client = FakeClient()
    target = tmp_path / "3.pdf"
    assert await client.download_media(make_message(CHAT_ID, 3, "photo"), file=target) is None
    assert not target.exists()


async def test_fake_download_media_raises_a_registered_error(tmp_path: Path) -> None:
    client = FakeClient(downloads={(CHAT_ID, 3): OSError("disk full")})
    with pytest.raises(OSError, match="disk full"):
        await client.download_media(make_message(CHAT_ID, 3, "photo"), file=tmp_path / "3.pdf")
